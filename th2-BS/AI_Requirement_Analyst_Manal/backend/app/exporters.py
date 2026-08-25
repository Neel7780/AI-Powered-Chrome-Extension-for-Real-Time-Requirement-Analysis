from datetime import datetime
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

def stamp() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S")

def report_text(data) -> str:
    lines = [
        "AI-POWERED REQUIREMENT ANALYSIS REPORT",
        "=" * 56,
        f"Generated on: {datetime.now():%Y-%m-%d %H:%M:%S}",
        "",
        "1. BASELINE REQUIREMENTS",
        "-" * 36,
        data.baseline or "No baseline requirements available.",
        "",
        "2. CLARIFICATION QUESTIONS",
        "-" * 36,
    ]
    if data.questions:
        lines.extend(f"Q{i}: {q}" for i, q in enumerate(data.questions, 1))
    else:
        lines.append("No clarification questions available.")
    lines += ["", "3. STAKEHOLDER RESPONSES", "-" * 36]
    if data.responses:
        lines.extend(f"Response {i}: {r}" for i, r in enumerate(data.responses, 1))
    else:
        lines.append("No stakeholder responses available.")
    lines += [
        "", "4. REFINED REQUIREMENTS", "-" * 36,
        data.refined or "No refined requirements available.",
        "", "5. QUALITY EVALUATION", "-" * 36,
        data.evaluation or "No quality evaluation available.",
        "", "6. CONCLUSION", "-" * 36,
        "AI-assisted clarification was used to reduce ambiguity and produce "
        "more specific and testable requirements.",
    ]
    return "\n".join(lines)

def write_txt(data, folder: Path) -> Path:
    path = folder / f"requirements_report_{stamp()}.txt"
    path.write_text(report_text(data), encoding="utf-8")
    return path

def write_pdf(data, folder: Path) -> Path:
    path = folder / f"requirements_report_{stamp()}.pdf"
    doc = SimpleDocTemplate(
        str(path), pagesize=A4, rightMargin=42, leftMargin=42,
        topMargin=42, bottomMargin=42
    )
    styles = getSampleStyleSheet()
    story = [
        Paragraph("AI-Powered Requirement Analysis Report", styles["Title"]),
        Spacer(1, 12),
        Paragraph(escape(f"Generated on: {datetime.now():%Y-%m-%d %H:%M:%S}"),
                  styles["BodyText"]),
        Spacer(1, 18),
    ]
    sections = [
        ("1. Baseline Requirements", data.baseline),
        ("2. Clarification Questions",
         "\n".join(f"Q{i}: {q}" for i, q in enumerate(data.questions, 1))),
        ("3. Stakeholder Responses",
         "\n".join(f"Response {i}: {r}" for i, r in enumerate(data.responses, 1))),
        ("4. Refined Requirements", data.refined),
        ("5. Quality Evaluation", data.evaluation),
        ("6. Conclusion",
         "AI-assisted clarification was used to reduce ambiguity and produce "
         "more specific and testable requirements."),
    ]
    for title, content in sections:
        story.append(Paragraph(escape(title), styles["Heading2"]))
        story.append(Spacer(1, 6))
        safe = escape(content or "No information available.").replace("\n", "<br/>")
        story.append(Paragraph(safe, styles["BodyText"]))
        story.append(Spacer(1, 14))
    doc.build(story)
    return path

def write_docx(data, folder: Path) -> Path:
    path = folder / f"requirements_report_{stamp()}.docx"
    doc = Document()
    doc.add_heading("AI-Powered Requirement Analysis Report", 0)
    doc.add_paragraph(f"Generated on: {datetime.now():%Y-%m-%d %H:%M:%S}")
    sections = [
        ("1. Baseline Requirements", data.baseline),
        ("2. Clarification Questions",
         "\n".join(f"Q{i}: {q}" for i, q in enumerate(data.questions, 1))),
        ("3. Stakeholder Responses",
         "\n".join(f"Response {i}: {r}" for i, r in enumerate(data.responses, 1))),
        ("4. Refined Requirements", data.refined),
        ("5. Quality Evaluation", data.evaluation),
        ("6. Conclusion",
         "AI-assisted clarification was used to reduce ambiguity and produce "
         "more specific and testable requirements."),
    ]
    for title, content in sections:
        doc.add_heading(title, 1)
        for paragraph in (content or "No information available.").split("\n"):
            doc.add_paragraph(paragraph)
    doc.save(path)
    return path
